import os
from docx import Document
from autogen_ext.models.openai import AzureOpenAIChatCompletionClient
from dotenv import load_dotenv

load_dotenv()

my_variable = os.getenv('MY_VARIABLE')

model_client = AzureOpenAIChatCompletionClient(
    api_version="2024-02-15-preview",
    azure_endpoint=os.environ['AZURE_ENDPOINT'],
    api_key=os.environ['API_KEY'],
    azure_deployment="gpt-4o",
    model="gpt-4o-2024-05-13"
)


def create_document(data, folder, file_name):
    print("Creating vacation document for user")
    doc = Document()
    doc.add_heading("Vacation Plans", level=1)
    doc.add_heading("User Request", level=2)
    doc.add_paragraph(data.request)
    doc.add_heading("Vacation Specification", level=2)
    doc.add_paragraph(data.specification)
    doc.add_heading("Vacation Budget", level=2)
    doc.add_paragraph(data.budget)
    doc.add_heading("Vacation Itinerary", level=2)
    doc.add_paragraph(data.itinerary)
    file_path = folder + '/' + file_name
    doc.save(file_path)
    os.system(f'start {file_path}')
